import csv
import io

from openpyxl import load_workbook


def _key(header):
    return "_".join(str(header).strip().lower().split())


def _clean(value):
    return value.strip() if isinstance(value, str) else value


def read_rows(file):
    """Rows of an uploaded .xlsx (first sheet) or .csv as dicts keyed by normalised header.

    Headers are lowercased with spaces turned into underscores, strings are stripped,
    blank rows are dropped, and each row carries its spreadsheet row number in "_row".
    """
    name = (getattr(file, "name", "") or "").lower()
    if not name.endswith((".xlsx", ".csv")):
        raise ValueError("Upload an .xlsx or .csv file.")
    file.seek(0)
    try:
        if name.endswith(".xlsx"):
            sheet = load_workbook(io.BytesIO(file.read()), read_only=True, data_only=True).worksheets[0]
            raw = list(sheet.iter_rows(values_only=True))
        else:
            raw = list(csv.reader(io.StringIO(file.read().decode("utf-8-sig"))))
    except Exception as exc:  # corrupt, renamed or wrongly encoded files
        raise ValueError(
            "Could not read the file. Save it as an Excel workbook (.xlsx) or a UTF-8 .csv and try again."
        ) from exc

    if not raw:
        return []
    headers = [_key(h) if h is not None and str(h).strip() else None for h in raw[0]]
    rows = []
    for number, values in enumerate(raw[1:], start=2):
        row = {h: _clean(v) for h, v in zip(headers, values) if h}
        if all(v is None or v == "" for v in row.values()):
            continue
        for h in headers:
            if h and h not in row:
                row[h] = ""
        row["_row"] = number
        rows.append(row)
    return rows
