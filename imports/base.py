"""The parts every spreadsheet import shares: fields, column matching and row plans."""

import io
import re
from dataclasses import dataclass, field


def norm(text):
    """'Faculty  Name ' / 'faculty_name' / 'FACULTY-NAME' -> 'faculty name'."""
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


@dataclass(frozen=True)
class Field:
    key: str
    label: str
    required: bool = False
    synonyms: tuple = ()
    help: str = ""

    def names(self):
        return {norm(self.key), norm(self.label), *(norm(s) for s in self.synonyms)}


@dataclass(frozen=True)
class Option:
    """A setting chosen on the column page (e.g. the department for rows without one)."""

    name: str
    label: str
    kind: str = "bool"  # bool | text | select
    choices: tuple = ()  # (value, label) for select
    default: object = None
    help: str = ""


class Status:
    NEW = "new"
    UPDATE = "update"
    SAME = "same"
    REVIEW = "review"
    INVALID = "invalid"

    LABELS = {
        NEW: "New",
        UPDATE: "Update",
        SAME: "No change",
        REVIEW: "Needs your decision",
        INVALID: "Needs correction",
    }


@dataclass
class RowPlan:
    number: int
    values: dict
    label: str = ""
    status: str = Status.NEW
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    match: object = None
    candidates: list = field(default_factory=list)
    changes: list = field(default_factory=list)  # (field label, old, new)

    @property
    def status_label(self):
        return Status.LABELS[self.status]

    def finish(self):
        if self.errors:
            self.status = Status.INVALID
        return self


@dataclass
class Plan:
    rows: list
    notes: list = field(default_factory=list)  # whole-file messages, e.g. categories to be added
    blocking: str = ""  # when set, the import can't be confirmed at all

    def count(self, status):
        return sum(1 for r in self.rows if r.status == status)

    @property
    def counts(self):
        return {s: self.count(s) for s in Status.LABELS}


class ImportKind:
    key = ""
    title = ""
    description = ""
    fields: tuple = ()
    options: tuple = ()
    modes = (("update", "Update existing records and add new ones"), ("new_only", "Only add new records; leave existing ones unchanged"))
    needs_period = False
    selectable = True  # rows can be ticked individually (not for a schedule, which is all or nothing)

    def guess_mapping(self, headers):
        """{field key: column index} from the file's headings, exact names first, then partial."""
        mapping, used = {}, set()
        normalised = [norm(h) for h in headers]
        for f in self.fields:
            for i, h in enumerate(normalised):
                if i not in used and h in f.names():
                    mapping[f.key] = i
                    used.add(i)
                    break
        for f in self.fields:
            if f.key in mapping:
                continue
            for i, h in enumerate(normalised):
                if i not in used and h and any(n and (n in h or h in n) for n in f.names() if len(n) > 3):
                    mapping[f.key] = i
                    used.add(i)
                    break
        return mapping

    def default_options(self):
        return {o.name: o.default for o in self.options}

    def values(self, batch, cells):
        """{field key: cleaned string} for one row, using the batch's column mapping."""
        out = {}
        for f in self.fields:
            index = batch.mapping.get(f.key)
            value = cells[index] if index is not None and index < len(cells) else None
            out[f.key] = "" if value is None else str(value).strip()
        return out

    def plan(self, batch) -> Plan:
        raise NotImplementedError

    def commit(self, batch, selected, decisions, mode, actor) -> dict:
        raise NotImplementedError


def parse_count(text, minimum=0):
    """'40' / '40.0' -> 40; None for blank; raises ValueError for anything else."""
    if text == "":
        return None
    try:
        number = float(text)
    except ValueError:
        raise ValueError from None
    if not number.is_integer() or number < minimum:
        raise ValueError
    return int(number)


def error_report(batch, plan):
    """An .xlsx with every row of the file plus its status and problems."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Import check"
    ws.append(["Row", "Status", "Problems", "Notes", *batch.headers])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    cells_by_row = {number: cells for number, cells in batch.rows}
    red = PatternFill("solid", fgColor="FDE2E1")
    for p in plan.rows:
        ws.append([p.number, p.status_label, "; ".join(p.errors), "; ".join(p.warnings), *cells_by_row.get(p.number, [])])
        if p.errors:
            for cell in ws[ws.max_row][:3]:
                cell.fill = red
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 40
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def template_workbook(kind):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = kind.title[:30]
    ws.append([f.label for f in kind.fields])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = max(14, len(cell.value) + 4)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
