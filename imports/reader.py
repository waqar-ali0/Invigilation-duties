"""Read an uploaded spreadsheet into headers and rows of plain JSON-safe values."""

import csv
import io
from datetime import date, datetime, time

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000


def _plain(value):
    """Spreadsheet cell -> str / int / float / None that can be stored as JSON."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d") if value.time() == time(0) else value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    if isinstance(value, int):
        return value
    text = str(value).strip()
    return text or None


def _xls_rows(data):
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
            else:
                row.append(cell.value if cell.value != "" else None)
        rows.append(row)
    return rows


def _csv_rows(data):
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    return list(csv.reader(io.StringIO(text)))


def read_table(file):
    """(headers, rows) of the first sheet of an .xlsx / .xls / .csv upload.

    `rows` is [[spreadsheet row number, [values...]], ...] with blank rows dropped. Headers keep
    their original wording (blank ones become "Column C"); repeated headers get " (2)".
    Raises ValueError with a message for the user when the file can't be used.
    """
    name = (getattr(file, "name", "") or "").lower()
    if not name.endswith((".xlsx", ".xlsm", ".xls", ".csv")):
        raise ValueError("Upload an Excel workbook (.xlsx or .xls) or a .csv file.")
    if getattr(file, "size", 0) > MAX_BYTES:
        raise ValueError("The file is larger than 10 MB.")
    file.seek(0)
    data = file.read()
    try:
        if name.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook

            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            raw = [list(r) for r in sheet.iter_rows(values_only=True)]
        elif name.endswith(".xls"):
            raw = _xls_rows(data)
        else:
            raw = _csv_rows(data)
    except Exception as exc:  # corrupt, renamed or wrongly encoded files
        raise ValueError("Could not read the file. Open it in Excel, save it again and retry.") from exc

    skipped = 0
    while raw and all(_plain(v) is None for v in raw[0]):
        raw.pop(0)  # blank lines above the headings
        skipped += 1
    if not raw:
        raise ValueError("The file is empty.")
    width = max(len(r) for r in raw)
    headers, seen = [], {}
    for i in range(width):
        value = _plain(raw[0][i]) if i < len(raw[0]) else None
        label = str(value) if value is not None else f"Column {_column_letter(i)}"
        seen[label] = seen.get(label, 0) + 1
        headers.append(label if seen[label] == 1 else f"{label} ({seen[label]})")
    rows = []
    for number, values in enumerate(raw[1:], start=2 + skipped):
        cells = [_plain(values[i]) if i < len(values) else None for i in range(width)]
        if all(v is None for v in cells):
            continue
        rows.append([number, cells])
    if len(rows) > MAX_ROWS:
        raise ValueError(f"The file has more than {MAX_ROWS} rows. Split it into smaller files.")
    # Columns that are blank in every row and have no heading are noise.
    keep = [i for i in range(width) if not headers[i].startswith("Column ") or any(r[1][i] is not None for r in rows)]
    return [headers[i] for i in keep], [[n, [cells[i] for i in keep]] for n, cells in rows]


def _column_letter(index):
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters
