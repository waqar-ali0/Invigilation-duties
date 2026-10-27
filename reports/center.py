"""The report centre: every report is a title, columns and rows built from the same records and
the same workload calculator as the dashboards, so the numbers always agree."""

import csv
import io
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime

from django.db.models import Q
from django.utils import timezone

from audit.models import AuditLog
from duties.models import (
    BackupMember,
    BackupReservation,
    Duty,
    DutyStatus,
    Obligation,
    Preference,
    PreferenceKind,
    ReservationStatus,
)
from exams.models import ExamPeriod
from workload.calc import Flag, period_loads
from workload.models import FacultyRule
from workload.rules import courses_taught


@dataclass
class Report:
    title: str
    columns: list
    rows: list = field(default_factory=list)
    notes: list = field(default_factory=list)


@dataclass
class Filters:
    period: ExamPeriod | None = None
    compare_with: ExamPeriod | None = None
    date_from: date | None = None
    date_to: date | None = None
    department: int | None = None
    category: int | None = None
    faculty: int | None = None
    room: str = ""
    building: int | None = None
    status: str = ""

    @classmethod
    def from_query(cls, q):
        def number(name):
            v = q.get(name, "")
            return int(v) if v.isdigit() else None

        def day(name):
            try:
                return date.fromisoformat(q.get(name, ""))
            except ValueError:
                return None

        period = ExamPeriod.objects.filter(pk=number("period")).first() if number("period") else None
        other = ExamPeriod.objects.filter(pk=number("compare_with")).first() if number("compare_with") else None
        status = q.get("status", "")
        return cls(period, other, day("date_from"), day("date_to"), number("department"), number("category"),
                   number("faculty"), q.get("room", "").strip(), number("building"), status if status in DutyStatus.values else "")

    def people_ok(self, f):
        return (
            (self.department is None or f.department_id == self.department)
            and (self.category is None or f.category_id == self.category)
            and (self.faculty is None or f.pk == self.faculty)
        )

    def duty_q(self, prefix=""):
        """Q() limiting duties (field names under `prefix`) to the filters."""
        p = prefix
        q = Q()
        if self.period:
            q &= Q(**{f"{p}period": self.period})
        if self.date_from:
            q &= Q(**{f"{p}room_allocation__session__slot__date__gte": self.date_from})
        if self.date_to:
            q &= Q(**{f"{p}room_allocation__session__slot__date__lte": self.date_to})
        if self.department:
            q &= Q(**{f"{p}faculty__department_id": self.department})
        if self.category:
            q &= Q(**{f"{p}faculty__category_id": self.category})
        if self.faculty:
            q &= Q(**{f"{p}faculty_id": self.faculty})
        if self.room:
            q &= Q(**{f"{p}room_allocation__room__icontains": self.room})
        if self.building:
            q &= Q(**{f"{p}room_allocation__room_ref__building_id": self.building})
        return q


def _duty_rows(f, order):
    qs = (
        Duty.objects.filter(f.duty_q())
        .select_related("faculty__department", "faculty__category", "period", "room_allocation__session__slot",
                        "room_allocation__room_ref__building", "replaces__faculty")
    )
    if f.status:
        qs = qs.filter(status=f.status)
    qs = qs.order_by(*order)
    rows = []
    for d in qs:
        s = d.room_allocation.session
        decision = "" if d.status != DutyStatus.ABSENT else ("not decided" if d.absence_excused is None else ("excused" if d.absence_excused else "unexcused"))
        rows.append([
            d.faculty.employee_id, d.faculty.name, d.faculty.department.name, d.faculty.category.name, d.period.name,
            s.slot.date, f"{s.slot.start_time:%H:%M}–{s.slot.end_time:%H:%M}", s.course_code, d.room_allocation.room,
            d.room_allocation.room_ref.building.name if d.room_allocation.room_ref_id else "",
            d.get_source_display(), d.get_status_display(), decision,
            d.replaces.faculty.name if d.replaces_id else "",
        ])
    return rows


DUTY_COLUMNS = ["Faculty ID", "Name", "Department", "Designation", "Period", "Date", "Time", "Course", "Room", "Building",
                "How", "Status", "Absence", "Took over from"]
SLOT_ORDER = ["room_allocation__session__slot__date", "room_allocation__session__slot__start_time"]


def faculty_wise(f):
    return Report("Faculty-wise duty allocation", DUTY_COLUMNS, _duty_rows(f, ["faculty__name", *SLOT_ORDER]))


def room_wise(f):
    return Report("Room-wise examination schedule", DUTY_COLUMNS, _duty_rows(f, ["room_allocation__room", *SLOT_ORDER, "faculty__name"]))


def slot_wise(f):
    return Report("Date-wise and slot-wise duty list", DUTY_COLUMNS, _duty_rows(f, [*SLOT_ORDER, "room_allocation__room", "faculty__name"]))


def _loads(f):
    if not f.period:
        return []
    return [x for x in period_loads(f.period) if f.people_ok(x.faculty)]


LOAD_COLUMNS = ["Faculty ID", "Name", "Department", "Designation", "Rule", "Assigned", "Completed", "Pending", "Missed",
                "Handed over", "Counted", "Remaining", "Status"]


def _load_row(x):
    return [x.faculty.employee_id, x.faculty.name, x.faculty.department.name, x.faculty.category.name, x.rule.describe(),
            x.assigned, x.completed, x.pending, x.missed, x.cancelled, x.counted,
            "" if x.remaining is None else x.remaining, ", ".join(label for label, _ in x.flag_labels)]


def workload(f):
    return Report("Workload per person", LOAD_COLUMNS, [_load_row(x) for x in _loads(f)])


def _group(f, key_name, key):
    groups = defaultdict(list)
    for x in _loads(f):
        groups[key(x)].append(x)
    rows = []
    for name, items in sorted(groups.items()):
        n = len(items)
        counted = sum(x.counted for x in items)
        rows.append([name, n, sum(x.assigned for x in items), sum(x.completed for x in items), sum(x.pending for x in items),
                     sum(x.missed for x in items), sum(x.cancelled for x in items), counted, round(counted / n, 2) if n else 0,
                     min((x.counted for x in items), default=0), max((x.counted for x in items), default=0)])
    return Report(f"{key_name}-wise workload", [key_name, "Faculty", "Assigned", "Completed", "Pending", "Missed", "Handed over",
                                              "Counted", "Average", "Lowest", "Highest"], rows)


def department_workload(f):
    return _group(f, "Department", lambda x: x.faculty.department.name)


def designation_workload(f):
    return _group(f, "Designation", lambda x: x.faculty.category.name)


def visiting(f):
    if not f.period:
        return Report("Visiting faculty course-based allocation", [])
    taught = courses_taught(f.period)
    rows = []
    for x in _loads(f):
        if not x.faculty.category.is_visiting:
            continue
        rows.append([x.faculty.employee_id, x.faculty.name, x.faculty.department.name, taught.get(x.faculty.pk, 0),
                     x.rule.note or x.rule.describe(), "" if x.rule.target is None else x.rule.target, x.assigned, x.completed, x.counted,
                     "" if x.remaining is None else x.remaining])
    return Report("Visiting faculty course-based allocation",
                  ["Faculty ID", "Name", "Department", "Courses with exams", "Rule", "Target", "Assigned", "Completed", "Counted", "Remaining"], rows)


def assigned_completed(f):
    rows = []
    for x in _loads(f):
        rows.append([x.faculty.employee_id, x.faculty.name, x.faculty.department.name, x.assigned, x.completed, x.pending, x.missed,
                     f"{round(100 * x.completed / x.assigned)}%" if x.assigned else ""])
    return Report("Assigned versus completed duties", ["Faculty ID", "Name", "Department", "Assigned", "Completed", "Pending", "Missed", "Completed %"], rows,
                  ["Assigned duties are never counted as completed until attendance is marked present."])


def missed(f):
    qs = Duty.objects.filter(f.duty_q(), status=DutyStatus.ABSENT).select_related(
        "faculty__department", "period", "room_allocation__session__slot", "absence_decided_by", "obligation__makeup__room_allocation__session__slot"
    ).prefetch_related("replaced_by__faculty").order_by(*SLOT_ORDER, "faculty__name")
    rows = []
    for d in qs:
        s = d.room_allocation.session
        o = getattr(d, "obligation", None)
        rows.append([d.faculty.employee_id, d.faculty.name, d.faculty.department.name, d.period.name, s.slot.date,
                     f"{s.slot.start_time:%H:%M}", s.course_code, d.room_allocation.room, d.absence_reason,
                     "not decided" if d.absence_excused is None else ("excused" if d.absence_excused else "unexcused"),
                     ", ".join(r.faculty.name for r in d.replaced_by.all()), o.status_label if o else "none owed"])
    return Report("Missed duties", ["Faculty ID", "Name", "Department", "Period", "Date", "Start", "Course", "Room", "Reason",
                                    "Decision", "Covered by", "Make-up"], rows)


def _obligation_rows(f, outstanding_only):
    qs = Obligation.objects.filter(f.duty_q("original__")).select_related(
        "owed_by__department", "original__room_allocation__session__slot", "original__period", "makeup__room_allocation__session__slot", "makeup__faculty"
    )
    rows = []
    for o in qs:
        if outstanding_only and not o.is_outstanding:
            continue
        m = o.makeup
        rows.append([o.owed_by.employee_id, o.owed_by.name, o.owed_by.department.name, o.original.period.name,
                     o.original.room_allocation.session.slot.date, o.original.room_allocation.room, o.status_label,
                     m.faculty.name if m else "", m.room_allocation.session.slot.date if m else "", m.room_allocation.room if m else "", o.note])
    return rows


OBLIGATION_COLUMNS = ["Faculty ID", "Name", "Department", "Period", "Missed on", "Missed room", "Status", "Make-up by", "Make-up date", "Make-up room", "Note"]


def obligations(f):
    return Report("Replacement (make-up) obligations", OBLIGATION_COLUMNS, _obligation_rows(f, False))


def outstanding(f):
    return Report("Faculty with outstanding obligations", OBLIGATION_COLUMNS, _obligation_rows(f, True))


def backups(f):
    reservations = BackupReservation.objects.select_related("duty")
    if f.period:
        reservations = reservations.filter(period=f.period)
    by = defaultdict(Counter)
    for r in reservations:
        c = by[r.faculty_id]
        c["reserved"] += 1
        if r.status == ReservationStatus.CALLED:
            c["called"] += 1
            if r.duty_id:
                c["assigned"] += 1
                if r.duty.status == DutyStatus.COMPLETED:
                    c["performed"] += 1
    rows = []
    for m in BackupMember.objects.select_related("faculty__department", "faculty__category"):
        if not f.people_ok(m.faculty):
            continue
        c = by[m.faculty_id]
        rows.append([m.faculty.employee_id, m.faculty.name, m.faculty.category.name, m.get_role_display(), m.priority,
                     c["reserved"], c["called"], c["assigned"], c["performed"]])
    return Report("Backup faculty utilisation", ["Faculty ID", "Name", "Designation", "Role", "Priority", "Reserved", "Called", "Assigned", "Performed"], rows)


def unfilled(f):
    from allocation.services import unfilled_rooms

    if not f.period:
        return Report("Unfilled room assignments", [])
    rows = [[r.session.slot.date, f"{r.session.slot.start_time:%H:%M}", r.session.course_code, r.room, r.invigilators_required, m, reason]
            for r, m, reason in unfilled_rooms(f.period)]
    return Report("Unfilled room assignments", ["Date", "Start", "Course", "Room", "Needed", "Missing", "Reason"], rows)


def fairness(f):
    groups = defaultdict(list)
    for x in _loads(f):
        if Flag.EXEMPT not in x.flags and Flag.NOT_ELIGIBLE not in x.flags:
            groups[x.faculty.category.name].append(x.counted)
            groups["All eligible faculty"].append(x.counted)
    rows = []
    for name, counts in sorted(groups.items(), key=lambda kv: (kv[0] != "All eligible faculty", kv[0])):
        avg = sum(counts) / len(counts)
        spread = (sum((c - avg) ** 2 for c in counts) / len(counts)) ** 0.5
        rows.append([name, len(counts), min(counts), max(counts), round(avg, 2), round(spread, 2), max(counts) - min(counts)])
    return Report("Workload fairness", ["Group", "Faculty", "Lowest", "Highest", "Average", "Spread (std. dev.)", "Range"], rows,
                  ["Counted duties under the period's counting rules. Different targets are by design."])


def preferences(f):
    if not f.period:
        return Report("Faculty preference satisfaction", [])
    wishes = defaultdict(list)
    for p in Preference.objects.filter(period=f.period).select_related("faculty__department"):
        wishes[p.faculty].append(p)
    duties = defaultdict(list)
    for d in Duty.objects.active().filter(period=f.period).select_related("room_allocation__session__slot", "room_allocation__room_ref"):
        duties[d.faculty_id].append(d)
    rows = []
    for person, prefs in sorted(wishes.items(), key=lambda kv: kv[0].name):
        if not f.people_ok(person):
            continue
        dates = {p.date for p in prefs if p.kind == PreferenceKind.DATE}
        slots = {p.slot_id for p in prefs if p.kind == PreferenceKind.SLOT}
        avoid = {p.slot_id for p in prefs if p.kind == PreferenceKind.AVOID_SLOT}
        buildings = {p.building_id for p in prefs if p.kind == PreferenceKind.BUILDING}
        met = 0
        mine = duties[person.pk]
        for d in mine:
            slot = d.room_allocation.session.slot
            ok = slot.id not in avoid
            if dates or slots:
                ok = ok and (slot.date in dates or slot.id in slots)
            if buildings:
                ok = ok and d.room_allocation.room_ref_id is not None and d.room_allocation.room_ref.building_id in buildings
            met += ok
        rows.append([person.employee_id, person.name, person.department.name, len(prefs), len(mine), met,
                     f"{round(100 * met / len(mine))}%" if mine else ""])
    return Report("Faculty preference satisfaction", ["Faculty ID", "Name", "Department", "Preferences", "Duties", "Duties that suit", "Suit %"], rows)


def comparison(f):
    if not (f.period and f.compare_with):
        return Report("Period comparison", [], notes=["Choose two periods."])
    a = {x.faculty.pk: x for x in period_loads(f.period)}
    b = {x.faculty.pk: x for x in period_loads(f.compare_with)}
    rows = []
    for pk in sorted(a.keys() | b.keys(), key=lambda k: (a.get(k) or b.get(k)).faculty.name):
        x, y = a.get(pk), b.get(pk)
        person = (x or y).faculty
        if not f.people_ok(person):
            continue
        rows.append([person.employee_id, person.name, person.category.name,
                     x.counted if x else 0, x.rule.describe() if x else "", y.counted if y else 0, y.rule.describe() if y else "",
                     (y.counted if y else 0) - (x.counted if x else 0)])
    return Report(f"{f.period.name} compared with {f.compare_with.name}",
                  ["Faculty ID", "Name", "Designation", f"Counted in {f.period.name}", "Rule", f"Counted in {f.compare_with.name}", "Rule", "Difference"], rows,
                  ["Each period is counted under its own rules; neither is changed."])


def overrides(f):
    rows = []
    entries = AuditLog.objects.filter(action__in=["duty.assign", "duty.replace", "duty.move", "duty.stand_in"]).select_related("actor").order_by("-timestamp")
    for e in entries[:2000]:
        after = e.after or {}
        if e.action == "duty.assign" and not after.get("warnings"):
            continue
        detail = after.get("override_reason") or after.get("reason") or ""
        rows.append([e.timestamp, e.actor.email if e.actor else "system", e.action.replace("duty.", ""), "; ".join(after.get("warnings", [])), detail])
    for r in FacultyRule.objects.filter(Q(exempt=True) | Q(allow_inactive=True)).select_related("faculty", "period", "updated_by"):
        if f.period and r.period_id != f.period.pk:
            continue
        what = "exempt this period" if r.exempt else "inactive exception"
        rows.append([r.updated_at, r.updated_by.email if r.updated_by else "", what, f"{r.faculty.name} · {r.period.name}", r.reason])
    rows.sort(key=lambda row: row[0], reverse=True)
    return Report("Manual overrides and exceptions", ["When", "By", "What", "Rules / who", "Reason"], rows)


def over_limit(f):
    return Report("Faculty exceeding their approved workload", LOAD_COLUMNS, [_load_row(x) for x in _loads(f) if Flag.OVER in x.flags])


REPORTS = {
    "faculty": ("Faculty-wise duty allocation", faculty_wise),
    "rooms": ("Room-wise examination schedule", room_wise),
    "slots": ("Date-wise and slot-wise duty list", slot_wise),
    "departments": ("Department-wise workload", department_workload),
    "designations": ("Designation-wise workload", designation_workload),
    "visiting": ("Visiting faculty course-based allocation", visiting),
    "assigned": ("Assigned versus completed", assigned_completed),
    "missed": ("Missed duties", missed),
    "obligations": ("Make-up obligations", obligations),
    "backups": ("Backup faculty utilisation", backups),
    "unfilled": ("Unfilled room assignments", unfilled),
    "fairness": ("Workload fairness", fairness),
    "preferences": ("Preference satisfaction", preferences),
    "comparison": ("Period comparison", comparison),
    "overrides": ("Overrides and exceptions", overrides),
    "over": ("Exceeding approved workload", over_limit),
    "outstanding": ("Outstanding obligations", outstanding),
    "workload": ("Workload per person", workload),
}


def _cell(value):
    if isinstance(value, datetime):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M") if timezone.is_aware(value) else value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return "" if value is None else value


def to_csv(report):
    out = io.StringIO()
    out.write("﻿")  # Excel opens UTF-8 CSV correctly with a BOM
    w = csv.writer(out)
    w.writerow(report.columns)
    for row in report.rows:
        w.writerow([_cell(v) for v in row])
    return out.getvalue().encode("utf-8")


def to_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:30].replace("/", "-")
    ws.append(report.columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in report.rows:
        ws.append([_cell(v) if not isinstance(v, (int, float)) else v for v in row])
    for i, name in enumerate(report.columns, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, min(40, len(str(name)) + 4))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(report, subtitle=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7, leading=8.5)
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    data = [[Paragraph(f"<b>{c}</b>", small) for c in report.columns]]
    for row in report.rows:
        data.append([Paragraph(str(_cell(v)).replace("&", "&amp;").replace("<", "&lt;"), small) for v in row])
    story = [Paragraph(report.title, styles["Title"])]
    if subtitle:
        story.append(Paragraph(subtitle, styles["Normal"]))
    for note in report.notes:
        story.append(Paragraph(note, styles["Italic"]))
    story.append(Spacer(1, 6))
    if report.columns:
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f7")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
    if not report.rows:
        story.append(Paragraph("No rows.", styles["Normal"]))
    doc.build(story)
    return out.getvalue()
