import csv
import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from allocation.services import publish_period, run_allocation
from duties.models import Duty, DutyStatus
from reports.center import REPORTS
from workload.calc import period_loads

pytestmark = pytest.mark.django_db


@pytest.fixture
def allocated(small_period, admin_user):
    run_allocation(small_period, admin_user, seed=1)
    publish_period(small_period, admin_user, confirm_rules=True)
    return small_period


def test_dashboard_totals_match_the_records(admin_client, allocated):
    first = Duty.objects.filter(period=allocated).first()
    Duty.objects.filter(pk=first.pk).update(status=DutyStatus.COMPLETED)
    r = admin_client.get(reverse("dashboard"), {"period": allocated.pk})
    c = r.context
    duties = Duty.objects.filter(period=allocated)
    assert c["assigned"] == duties.exclude(status=DutyStatus.CANCELLED).count()
    assert (c["completed"], c["pending"], c["missed"]) == (1, duties.filter(status=DutyStatus.ASSIGNED).count(), 0)
    assert c["seats_needed"] == 6 and c["seats_filled"] == 6 and c["coverage"] == 100
    loads = period_loads(allocated)
    assert c["no_target"] == sum(1 for x in loads if x.rule.target is None)
    assert r.content.decode().count("<svg class=\"chart\"") >= 3


def test_every_card_links_somewhere_that_opens(admin_client, allocated):
    page = admin_client.get(reverse("dashboard"), {"period": allocated.pk}).content.decode()
    import re

    main = page[page.index("<main"):]
    links = set(re.findall(r'<a class="card" href="([^"]+)"', main))
    assert len(links) >= 8
    for href in links:
        assert admin_client.get(href.replace("&amp;", "&")).status_code == 200, href


def test_period_choice_is_remembered_and_shown_in_the_top_bar(admin_client, allocated):
    r = admin_client.get(reverse("pick_period"), {"period": allocated.pk, "next": "/dashboard/"})
    assert r.status_code == 302 and r.url == "/dashboard/"
    page = admin_client.get(reverse("dashboard")).content.decode()
    assert f'<option value="{allocated.pk}" selected>' in page


@pytest.mark.parametrize("key", list(REPORTS))
def test_every_report_renders_and_downloads(admin_client, allocated, key):
    params = {"type": key, "period": allocated.pk}
    assert admin_client.get(reverse("report_center"), params).status_code == 200
    for fmt in ("csv", "xlsx", "pdf"):
        r = admin_client.get(reverse("report_center"), {**params, "format": fmt})
        assert r.status_code == 200 and r["Content-Disposition"].endswith(f'.{fmt}"')


def test_report_numbers_match_the_workload_calculator(admin_client, allocated):
    r = admin_client.get(reverse("report_center"), {"type": "workload", "period": allocated.pk, "format": "csv"})
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    header, body = rows[0], rows[1:]
    counted = {row[header.index("Faculty ID")]: int(row[header.index("Counted")]) for row in body}
    assert counted == {x.faculty.employee_id: x.counted for x in period_loads(allocated)}


def test_report_filters_apply(admin_client, allocated):
    someone = Duty.objects.filter(period=allocated).first().faculty
    r = admin_client.get(reverse("report_center"), {"type": "faculty", "period": allocated.pk, "faculty": someone.pk, "format": "xlsx"})
    ws = load_workbook(io.BytesIO(r.content)).active
    names = {row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert names == {someone.name}


def test_timetable_grid(admin_client, allocated):
    page = admin_client.get(reverse("timetable", args=[allocated.pk])).content.decode()
    assert "4/4 seats" in page and "2/2 seats" in page


def test_dashboard_and_reports_are_admin_only(client, faculty_user, allocated):
    client.force_login(faculty_user)
    for url in (reverse("dashboard"), reverse("report_center"), reverse("timetable", args=[allocated.pk])):
        assert client.get(url).status_code in (302, 403)


def test_dashboard_without_any_period(admin_client):
    assert "No exam periods yet" in admin_client.get(reverse("dashboard")).content.decode()
