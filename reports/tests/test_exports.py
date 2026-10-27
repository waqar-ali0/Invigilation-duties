import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from allocation.services import run_allocation
from reports.excel import load_xlsx, roster_xlsx
from reports.pdf import roster_pdf
from reports.queries import roster_rows

pytestmark = pytest.mark.django_db


@pytest.fixture
def allocated_period(small_period, admin_user):
    run_allocation(small_period, admin_user, seed=1)
    return small_period


def sheet_rows(data):
    return [list(r) for r in load_workbook(io.BytesIO(data)).active.iter_rows(values_only=True)]


def test_roster_rows_ordered_with_names(allocated_period):
    rows = roster_rows(allocated_period)
    assert [(r.course_code, r.room) for r in rows] == [("CS101", "A1"), ("CS102", "A2"), ("CS201", "B1")]
    assert all(len(r.invigilators) == 2 and r.invigilators == sorted(r.invigilators) for r in rows)


def test_roster_rows_query_count_constant(django_assert_max_num_queries, allocated_period):
    with django_assert_max_num_queries(6):
        roster_rows(allocated_period)


def test_roster_pdf_is_pdf(admin_client, allocated_period):
    r = admin_client.get(reverse("roster_pdf", args=[allocated_period.pk]))
    assert r["Content-Type"] == "application/pdf" and r.content.startswith(b"%PDF")
    assert 'filename="fall-2026-finals-roster.pdf"' in r["Content-Disposition"]


def test_roster_pdf_function_returns_bytes(allocated_period):
    assert roster_pdf(allocated_period).startswith(b"%PDF")


def test_roster_xlsx_has_header_and_one_row_per_room(allocated_period):
    rows = sheet_rows(roster_xlsx(allocated_period))
    assert rows[0] == [
        "Date",
        "Start",
        "End",
        "Course Code",
        "Course Name",
        "Department",
        "Room",
        "Students",
        "Invigilators",
    ]
    assert len(rows) == 4 and rows[1][3] == "CS101" and rows[1][8].count(", ") == 1


def test_load_xlsx_includes_zero_duty_faculty_and_summary_rows(allocated_period):
    rows = sheet_rows(load_xlsx(allocated_period))
    assert rows[0] == ["Faculty ID", "Name", "Department", "Category", "Duties"]
    body = rows[1:9]
    assert len(body) == 8 and sum(r[4] for r in body) == 6 and min(r[4] for r in body) == 0
    assert [r[0] for r in rows[-3:]] == ["Min", "Max", "Average"]
    assert rows[-3][4] == 0 and rows[-2][4] == 1 and rows[-1][4] == 0.75


def test_exports_admin_only(client, faculty_user, allocated_period):
    client.force_login(faculty_user)
    for name in ("roster_pdf", "roster_xlsx", "load_xlsx"):
        assert client.get(reverse(name, args=[allocated_period.pk])).status_code == 403


def test_xlsx_views_download(admin_client, allocated_period):
    for name, filename in (("roster_xlsx", "fall-2026-finals-roster.xlsx"), ("load_xlsx", "fall-2026-finals-load.xlsx")):
        r = admin_client.get(reverse(name, args=[allocated_period.pk]))
        assert r.status_code == 200 and f'filename="{filename}"' in r["Content-Disposition"]


def test_roster_pdf_keeps_markup_like_course_names(allocated_period, monkeypatch):
    import reports.pdf
    from exams.models import ExamSession

    texts = []

    class Spy(reports.pdf.Paragraph):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            texts.append(self.getPlainText())

    ExamSession.objects.filter(course_code="CS101").update(course_name="Data <Structures> & Algorithms")
    monkeypatch.setattr(reports.pdf, "Paragraph", Spy)
    roster_pdf(allocated_period)
    assert "CS101 — Data <Structures> & Algorithms" in texts
