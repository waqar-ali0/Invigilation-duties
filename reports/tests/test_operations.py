import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import Role, User
from audit.log import log
from duties.handover import transfer_duty
from duties.models import Duty, DutySource, DutyStatus, SwapRequest, SwapStatus
from exams.models import ExamPeriod, PeriodStatus
from faculty.models import Department
from reports.excel import operations_xlsx
from reports.operations import department_summary, faculty_records, handover_history

pytestmark = pytest.mark.django_db


@pytest.fixture
def chain(period, make_room, make_faculty, admin_user):
    """A's duty swapped to B, then B's new duty replaced by C; D present once and absent once."""
    a, b, c, d = (make_faculty(n) for n in ("Alpha", "Bravo", "Charlie", "Delta"))
    room = make_room(day=1, room="R1")
    first = Duty.objects.create(faculty=a, room_allocation=room, source=DutySource.AUTO)
    second = transfer_duty(first, b, DutySource.SWAP)
    SwapRequest.objects.create(
        duty=first, requested_by=a, requested_with=b, status=SwapStatus.APPROVED, message="Wedding", new_duty=second, decided_by=admin_user
    )
    third = transfer_duty(second, c, DutySource.REPLACEMENT)
    log(admin_user, "duty.replace", second, before={"faculty": b.pk}, after={"faculty": c.pk, "reason": "Sick", "duty": third.pk})
    Duty.objects.create(faculty=d, room_allocation=make_room(day=2, room="R2"), source=DutySource.AUTO, status=DutyStatus.COMPLETED)
    Duty.objects.create(faculty=d, room_allocation=make_room(day=3, room="R3"), source=DutySource.AUTO, status=DutyStatus.ABSENT)
    ExamPeriod.objects.filter(pk=period.pk).update(status=PeriodStatus.PUBLISHED)
    return a, b, c, d


def counts(record):
    return (record.assigned, record.completed, record.absent, record.unmarked, record.swapped_away, record.taken_over, record.replaced)


def test_faculty_records_follow_the_chain(chain, period):
    records = {r.faculty.name: r for r in faculty_records(period)}
    assert counts(records["Alpha"]) == (0, 0, 0, 0, 1, 0, 0)
    assert counts(records["Bravo"]) == (0, 0, 0, 0, 0, 1, 1)
    assert counts(records["Charlie"]) == (1, 0, 0, 1, 0, 1, 0)
    assert counts(records["Delta"]) == (2, 1, 1, 0, 0, 0, 0)
    assert [d.isoformat() for d in records["Delta"].absence_dates] == ["2026-12-03"]


def test_department_summary(chain, period, dept):
    physics = Department.objects.create(name="Physics")
    rows = department_summary(period)
    cs = next(r for r in rows if r["department"] == dept.name)
    assert (cs["assigned"], cs["completed"], cs["absent"], cs["swapped_away"], cs["taken_over"], cs["replaced"]) == (3, 1, 1, 1, 2, 1)
    assert department_summary(period, [physics]) == []


def test_handover_history(chain, period):
    rows = handover_history(period)
    assert [(r["type"], r["from"].name, r["to"].name, r["note"]) for r in rows] == [
        ("Replacement", "Bravo", "Charlie", "Sick"),
        ("Swap", "Alpha", "Bravo", "Wedding"),
    ]


def test_operations_xlsx(chain, period):
    wb = load_workbook(io.BytesIO(operations_xlsx(period)))
    assert wb.sheetnames == ["Faculty record", "Department summary", "Swaps and replacements"]
    assert [c.value for c in wb["Faculty record"][1]] == [
        "Name", "Department", "Assigned", "Completed", "Absent", "Not yet marked", "Swapped away", "Taken over", "Replaced", "Absence dates",
    ]


def test_report_page_admin_only(admin_client, client, chain, period):
    page = admin_client.get(reverse("operations_report", args=[period.pk])).content.decode()
    assert "Charlie" in page and "Sick" in page
    assert admin_client.get(reverse("operations_xlsx", args=[period.pk])).status_code == 200
    client.force_login(User.objects.create_user("x@uni.edu", "x"))
    assert client.get(reverse("operations_report", args=[period.pk])).status_code == 403


def test_hod_department_duties_shows_summary(client, chain, period, dept, make_member):
    hod = make_member("Head", role=Role.HOD)
    dept.hod = hod
    dept.save()
    hod.user.refresh_from_db()
    client.force_login(hod.user)
    page = client.get(reverse("department_duties")).content.decode()
    assert "Completed" in page and "Absent" in page
