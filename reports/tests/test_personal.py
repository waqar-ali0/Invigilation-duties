import io

import pytest
from openpyxl import load_workbook

from duties.models import Duty, DutySource
from exams.models import ExamPeriod, PeriodStatus
from reports.excel import department_xlsx
from reports.personal import duties_ics, duties_pdf, faculty_duties

pytestmark = pytest.mark.django_db


@pytest.fixture
def my_duties(period, make_room, make_member):
    me = make_member("Ali Raza")
    a = make_room(day=1, start="09:00", end="12:00", course="CS101", room="Hall A")
    b = make_room(day=3, start="14:00", end="17:00", course="MA200", room="Lab, 2")
    a.session.course_name = "Intro; Programming"
    a.session.save()
    for r in (b, a):
        Duty.objects.create(faculty=me, room_allocation=r, source=DutySource.MANUAL)
    ExamPeriod.objects.filter(pk=period.pk).update(status=PeriodStatus.PUBLISHED)
    period.refresh_from_db()
    return me


def test_faculty_duties_ordered_and_hidden_for_drafts(my_duties, period):
    assert [d.room_allocation.room for d in faculty_duties(my_duties, period)] == ["Hall A", "Lab, 2"]
    ExamPeriod.objects.filter(pk=period.pk).update(status=PeriodStatus.ALLOCATING)
    assert faculty_duties(my_duties) == []


def test_ics_events_in_utc_from_local_time(my_duties, period, settings):
    assert settings.TIME_ZONE == "Asia/Karachi"
    text = duties_ics(my_duties, faculty_duties(my_duties, period))
    assert text.startswith("BEGIN:VCALENDAR\r\nVERSION:2.0\r\n") and text.endswith("END:VCALENDAR\r\n")
    assert text.count("BEGIN:VEVENT") == 2
    assert "DTSTART:20261201T040000Z" in text and "DTEND:20261201T070000Z" in text  # 09:00–12:00 PKT
    assert "SUMMARY:Invigilation: CS101" in text
    assert "LOCATION:Lab\\, 2" in text and "DESCRIPTION:Intro\\; Programming" in text
    duty_id = faculty_duties(my_duties, period)[0].pk
    assert f"UID:duty-{duty_id}@invigilation" in text


def test_personal_pdf(my_duties, period):
    assert duties_pdf(my_duties, period).startswith(b"%PDF")


def test_department_xlsx(my_duties, period, dept):
    rows = [list(r) for r in load_workbook(io.BytesIO(department_xlsx(period, [dept]))).active.iter_rows(values_only=True)]
    assert rows[0] == ["Name", "Department", "Duties", "Dates and rooms"]
    ali = next(r for r in rows if r[0] == "Ali Raza")
    assert ali[2] == 2 and "Hall A" in ali[3]
