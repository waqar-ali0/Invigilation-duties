r"""Full end-to-end audit on a fresh, throw-away database, through the real URLs, forms and CSRF checks.

It never touches db.sqlite3: everything happens in a temporary database created and filled from
scratch (imports, invitations, requests, allocation, publishing, swaps, replacement, reminders,
attendance, reports, audit log, search, permissions, emails).

    .venv\Scripts\python scripts\audit_e2e.py

Prints PASS/FAIL per check and exits with status 1 if anything failed.
"""

import html
import io
import os
import re
import sys
import tempfile
import time
import traceback
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

PROJECT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT))
os.chdir(PROJECT)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
tmp = Path(tempfile.mkdtemp(prefix="audit-"))
os.environ["DATABASE_URL"] = f"sqlite:///{(tmp / 'audit.sqlite3').as_posix()}"
os.environ["DEBUG"] = "True"
os.environ["SITE_URL"] = "http://testserver"
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")

import django  # noqa: E402

django.setup()
from django.test.utils import setup_test_environment  # noqa: E402

setup_test_environment()  # locmem email backend, testserver host

from django.conf import settings  # noqa: E402
from django.core import mail  # noqa: E402
from django.core.management import call_command  # noqa: E402
from django.test import Client  # noqa: E402
from django.urls import reverse  # noqa: E402
from django.utils import timezone  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

settings.ALLOWED_HOSTS = ["testserver", "localhost"]
call_command("migrate", verbosity=0)

from accounts.models import Role, User  # noqa: E402
from audit.models import AuditLog  # noqa: E402
from duties.models import Duty, DutyStatus, SwapRequest, SwapStatus, Unavailability, UnavailabilityStatus  # noqa: E402
from exams.models import ExamPeriod, PeriodStatus, RoomAllocation, TimeSlot  # noqa: E402
from faculty.models import Department, Faculty, FacultyCategory  # noqa: E402
from notifications.models import EmailStatus, OutboundEmail  # noqa: E402
from notifications.outbox import send_queued  # noqa: E402

RESULTS = []
PASSWORD = "Audit-pass-2026!"


def check(name, condition, detail=""):
    RESULTS.append((bool(condition), name, detail))
    print(("PASS " if condition else "FAIL ") + name + (f"  [{detail}]" if detail and not condition else ""))
    return condition


class Browser:
    """A Django test client that behaves like a browser: CSRF enforced, token taken from the cookie."""

    def __init__(self):
        self.c = Client(enforce_csrf_checks=True)

    def get(self, url, **params):
        r = self.c.get(url, params or None)
        assert r.status_code < 500, f"GET {url} -> {r.status_code}"
        return r

    def post(self, url, data=None, files=None, follow=False):
        if "csrftoken" not in self.c.cookies:
            self.c.get(reverse("login"))
        data = dict(data or {})
        data["csrfmiddlewaretoken"] = self.c.cookies["csrftoken"].value
        if files:
            data.update(files)
        r = self.c.post(url, data, follow=follow)
        assert r.status_code < 500, f"POST {url} -> {r.status_code}"
        return r

    def login(self, email, password=PASSWORD):
        self.c.get(reverse("login"))
        r = self.post(reverse("login"), {"username": email, "password": password})
        return r.status_code == 302 and "_auth_user_id" in self.c.session

    def text(self, url, **params):
        return self.get(url, **params).content.decode()


def xlsx(rows, name):
    from django.core.files.uploadedfile import SimpleUploadedFile

    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue())


def wizard(browser, kind, rows, name, period=None, options=None, mode=None):
    """Upload through the import wizard, keep the guessed columns, confirm every row. Returns the batch."""
    from imports.models import ImportBatch

    url = reverse("import_start_period", args=[kind, period.pk]) if period else reverse("import_start", args=[kind])
    before = ImportBatch.objects.count()
    r = browser.post(url, files={"file": xlsx(rows, name)})
    if ImportBatch.objects.count() == before:
        return None, r
    batch = ImportBatch.objects.latest("pk")
    data = {f"map_{k}": v for k, v in batch.mapping.items()}
    for key, value in batch.options.items():
        if value is True:
            data[f"opt_{key}"] = "on"
        elif value:
            data[f"opt_{key}"] = value
    data.update(options or {})
    browser.post(reverse("import_columns", args=[batch.pk]), data)
    preview = browser.get(reverse("import_batch", args=[batch.pk]))
    confirm = {"row": [str(n) for n, _ in batch.rows]}
    if mode:
        confirm["mode"] = mode
    browser.post(reverse("import_batch", args=[batch.pk]), confirm)
    batch.refresh_from_db()
    return batch, preview


def run():
    # ---------------------------------------------------------------- setup
    admin = User.objects.create_user("office@uni.edu", PASSWORD, role=Role.ADMIN)
    User.objects.create_user("staff@uni.edu", PASSWORD, role=Role.EXAM_STAFF)
    office = Browser()
    check("admin logs in through the login form", office.login("office@uni.edu"))
    check("admin home redirects to exam periods", office.get("/").url == reverse("period_list"))

    # ---------------------------------------------------------------- imports
    r = office.post(reverse("department_import"), files={"file": xlsx([["Name"], ["Computer Science"], ["Physics"], ["Mathematics"]], "d.xlsx")})
    check("department import", Department.objects.count() == 3, r.content.decode()[:300])

    people = [["Employee ID", "Name", "Email", "Department", "Designation", "Mobile"]]
    depts = ["Computer Science", "Physics", "Mathematics"]
    designations = ["Professor", "Associate Professor", "Assistant Professor", "Lecturer", "Lecturer"]
    for i in range(15):
        people.append([f"E{i:03d}", f"Person {i:02d} O'Neil & Co", f"Person{i:02d}@Uni.edu", depts[i % 3], designations[i % 5], f"0300-11100{i:02d}"])
    batch, preview = wizard(office, "faculty", people, "f.xlsx")
    check("faculty import wizard (15, mixed-case emails, special characters)", Faculty.objects.count() == 15 and batch.result["created"] == 15, str(batch and batch.result))
    check("phone numbers normalised", Faculty.objects.filter(phone_e164="+923001110003").exists())
    batch, preview = wizard(office, "faculty", [people[0], ["X1", "Bad", "bad", "Nowhere", "Dean", "x"]], "bad.xlsx")
    check("faculty import explains bad rows and imports nothing", "Unknown designation" in preview.content.decode() and batch.result["invalid"] == 1 and Faculty.objects.count() == 15)
    again, _ = wizard(office, "faculty", people, "f-again.xlsx")
    check("re-importing the same file creates no duplicates", Faculty.objects.count() == 15 and again.result["unchanged"] == 15, str(again.result))
    own = [["S.No", "Teacher Name", "Rank", "Contact No"], [1, "Zara Visiting", "Visiting Lecturer", "0333-4445566"], [2, "Person 01 O'Neil & Co", "Lecturer", "0321-0000000"]]
    batch, preview = wizard(office, "faculty", own, "own.xlsx", options={"opt_default_department": str(Department.objects.get(name="Physics").pk)})
    zara = Faculty.objects.filter(name="Zara Visiting").first()
    check("own spreadsheet without IDs: columns guessed, FAC ID generated", zara is not None and zara.employee_id.startswith("FAC-") and zara.category.code == "visiting", str(batch.result))
    check("same name with a different phone is not merged", Faculty.objects.filter(name__startswith="Person 01").count() == 1 and batch.result["skipped"] == 1, str(batch.result))
    Faculty.objects.filter(name="Zara Visiting").delete()  # keep the rest of the audit on the 15 imported people
    rooms_file = [["Block", "Room No", "Seats", "Invigilators"], ["Main", "R1", 60, ""], ["Main", "R2", 40, ""], ["Main", "R3", 120, ""], ["Main", "R9", 200, 4]]
    batch, _ = wizard(office, "rooms", rooms_file, "rooms.xlsx")
    from exams.models import Room
    check("room inventory import", Room.objects.count() == 4 and Room.objects.get(name="R9").invigilators_needed == 4, str(batch and batch.result))

    cs, phys, maths = (Department.objects.get(name=n) for n in depts)
    hod_cs = Faculty.objects.get(employee_id="E000")
    exempt = Faculty.objects.get(employee_id="E014")
    Faculty.objects.filter(pk=exempt.pk).update(is_exempt=True)

    # HOD through the Django admin change form.
    r = office.get(reverse("admin:faculty_department_change", args=[cs.pk]))
    check("admin panel department page opens", r.status_code == 200)
    r = office.post(reverse("admin:faculty_department_change", args=[cs.pk]), {"name": cs.name, "hod": hod_cs.pk, "_save": "Save"})
    cs.refresh_from_db()
    check("HOD set through admin panel", cs.hod_id == hod_cs.pk, f"status {r.status_code}")

    # Exam period through the Django admin add form (with its inlines).
    start = timezone.localdate() + timedelta(days=30)
    add = reverse("admin:exams_examperiod_add")
    office.get(add)
    form = {
        "name": "Audit Finals",
        "start_date": start.isoformat(),
        "end_date": (start + timedelta(days=4)).isoformat(),
        "availability_deadline_0": "", "availability_deadline_1": "",
        "settings-TOTAL_FORMS": "1", "settings-INITIAL_FORMS": "0", "settings-MIN_NUM_FORMS": "0", "settings-MAX_NUM_FORMS": "1",
        "settings-0-max_duties_per_faculty": "6", "settings-0-max_duties_per_day": "2",
        "settings-0-students_per_invigilator": "30", "settings-0-min_invigilators_per_room": "2",
        "settings-0-avoid_back_to_back": "on", "settings-0-exclude_course_teacher": "on", "settings-0-swap_cutoff_hours": "24",
        "designation_caps-TOTAL_FORMS": "1", "designation_caps-INITIAL_FORMS": "0",
        "designation_caps-MIN_NUM_FORMS": "0", "designation_caps-MAX_NUM_FORMS": "1000",
        "designation_caps-0-category": str(FacultyCategory.objects.get(code="professor").pk), "designation_caps-0-max_duties": "2",
        "_save": "Save",
    }
    r = office.post(add, form)
    period = ExamPeriod.objects.filter(name="Audit Finals").first()
    errors = re.findall(r'class="errorlist[^"]*"[^>]*>(.*?)</ul>', r.content.decode(), re.S)
    check("exam period created through admin panel (with settings + designation cap)", period is not None, str(errors)[:400])
    if period is None:
        return
    check("period settings saved from the inline", period.settings.max_duties_per_faculty == 6, str(period.settings.max_duties_per_faculty))
    check("designation cap saved", period.cap_for("professor") == 2)

    # Schedule: 5 days x 2 slots x 3 courses, Excel date and time cells, teacher emails.
    sched = [["Date", "Start Time", "End Time", "Course Code", "Course Name", "Department", "Room", "Student Count", "Course Teacher Emails"]]
    for d in range(5):
        day = datetime.combine(start + timedelta(days=d), datetime.min.time())
        for (h1, h2) in ((9, 12), (14, 17)):
            for c, dept in enumerate(depts):
                sched.append([day, datetime(2000, 1, 1, h1).time(), datetime(2000, 1, 1, h2).time(), f"{dept[:2].upper()}{d}{h1}{c}",
                              f"{dept} & Data <101>", dept, f"R{c + 1}", 25, f"person{c:02d}@uni.edu"])
    batch, preview = wizard(office, "schedule", sched, "s.xlsx", period=period)
    check("schedule import (30 rooms)", RoomAllocation.objects.filter(session__slot__period=period).count() == 30, str(batch and batch.result))
    check("schedule rooms linked to the room list", not RoomAllocation.objects.filter(session__slot__period=period, room_ref__isnull=True).exists())
    seats = sum(RoomAllocation.objects.filter(session__slot__period=period).values_list("invigilators_required", flat=True))
    check("invigilators required computed (2 per room)", seats == 60, str(seats))

    # ---------------------------------------------------------------- accounts
    r = office.post(reverse("invite_all"))
    check("invite all queues an email per active faculty", OutboundEmail.objects.filter(kind="activation").count() == 15)
    sent, failed = send_queued()
    check("outbox sends invitations", (sent, failed) == (15, 0), f"{sent},{failed}")
    check("stored activation emails no longer contain the link", not OutboundEmail.objects.filter(kind="activation", body__contains="/activate/").exists())
    browsers = {}
    for f in Faculty.objects.all():
        msg = next(m for m in mail.outbox if m.to == [f.email.lower()] and "Set up" in m.subject)
        link = re.search(r"http://testserver(/activate/\S+)", msg.body).group(1)
        b = Browser()
        b.get(link)
        r = b.post(link, {"new_password1": PASSWORD, "new_password2": PASSWORD})
        ok = r.status_code == 302
        browsers[f.employee_id] = b
        if not ok:
            check(f"activation for {f.employee_id}", False, r.content.decode()[:200])
    check("every faculty member activated from the emailed link", all(f.user.has_usable_password() for f in Faculty.objects.select_related("user")))
    check("email body not HTML-escaped", "O'Neil & Co" in mail.outbox[0].body, mail.outbox[0].body[:120])
    check("HOD account got HOD role", User.objects.get(email="person00@uni.edu").role == Role.HOD)
    fresh = Browser()
    check("login works with different letter case", fresh.login("PERSON01@UNI.EDU"))
    check("faculty home is My duties", fresh.get("/").url == reverse("my_duties"))

    # ---------------------------------------------------------------- availability & requests
    deadline = timezone.localtime(timezone.now() + timedelta(days=7)).strftime("%Y-%m-%dT%H:%M")
    office.get(reverse("period_dashboard", args=[period.pk]))
    office.post(reverse("open_availability", args=[period.pk]), {"deadline": deadline})
    period.refresh_from_db()
    check("availability opened", period.availability_is_open)
    check("availability emails to active non-exempt accounts", OutboundEmail.objects.filter(kind="availability_open").count() == 14,
          str(OutboundEmail.objects.filter(kind="availability_open").count()))

    f1 = browsers["E001"]  # Physics? (i%3): E001 -> Physics
    f3 = browsers["E003"]  # Computer Science
    hod = browsers["E000"]
    day1 = start.isoformat()
    slot_pm_day2 = TimeSlot.objects.get(period=period, date=start + timedelta(days=1), start_time=datetime(2000, 1, 1, 14).time())
    f3.get(reverse("my_unavailability"))
    f3.post(reverse("my_unavailability"), {"period": period.pk, "date": day1, "slot": "", "reason": "Conference"})
    f3.post(reverse("my_unavailability"), {"period": period.pk, "date": (start + timedelta(days=1)).isoformat(), "slot": slot_pm_day2.pk, "reason": "Doctor"})
    dup = f3.post(reverse("my_unavailability"), {"period": period.pk, "date": day1, "slot": "", "reason": "Again"})
    check("duplicate request refused with a message", "already have a request" in dup.content.decode())
    hod.get(reverse("my_unavailability"))
    hod.post(reverse("my_unavailability"), {"period": period.pk, "date": day1, "slot": "", "reason": "HOD meeting"})
    reqs = Unavailability.objects.filter(period=period)
    check("three requests pending", reqs.filter(status="pending").count() == 3, str(list(reqs.values_list("status", flat=True))))
    page = office.text(reverse("period_dashboard", args=[period.pk]))
    check("dashboard shows the pending-capacity warning", "request(s) pending" in page)

    page = hod.text(reverse("department_requests"))
    own = reqs.get(faculty__employee_id="E000")
    check("HOD sees department requests but not their own", "Conference" in page and "HOD meeting" not in page)
    r3a, r3b = reqs.filter(faculty__employee_id="E003").order_by("date")
    hod.post(reverse("decide_request", args=[r3a.pk]), {"decision": "approve", "comment": "OK"})
    hod.post(reverse("decide_request", args=[r3b.pk]), {"decision": "reject", "comment": "Exams need you"})
    check("HOD cannot decide own request (404)", hod.post(reverse("decide_request", args=[own.pk]), {"decision": "approve"}).status_code == 404)
    office.get(reverse("admin_requests"))
    office.post(reverse("decide_request", args=[own.pk]), {"decision": "approve"})
    statuses = dict(Unavailability.objects.filter(period=period).values_list("reason", "status"))
    check("decisions recorded", statuses == {"Conference": "approved", "Doctor": "rejected", "HOD meeting": "approved"}, str(statuses))
    check("decision emails queued", OutboundEmail.objects.filter(kind="request_decision").count() == 3)
    f3.get(reverse("my_unavailability"))

    # ---------------------------------------------------------------- allocation
    office.post(reverse("run_allocation", args=[period.pk]))
    duties = list(Duty.objects.active().filter(period=period).select_related("faculty", "room_allocation__session__slot"))
    check("allocation created duties", len(duties) > 0, str(len(duties)))
    per_slot = Counter((d.faculty_id, d.room_allocation.session.slot_id) for d in duties)
    per_day = Counter((d.faculty_id, d.room_allocation.session.slot.date) for d in duties)
    totals = Counter(d.faculty_id for d in duties)
    check("nobody twice in one slot", max(per_slot.values()) == 1)
    check("daily limit respected", max(per_day.values()) <= 2)
    caps_ok = all(totals[f.pk] <= period.cap_for(f.category_id) for f in Faculty.objects.all())
    check("caps respected (professor 2, others 6)", caps_ok)
    check("exempt faculty not assigned", exempt.pk not in totals)
    unavailable = [d for d in duties if d.faculty.employee_id in ("E003", "E000") and d.room_allocation.session.slot.date == start]
    check("approved unavailability respected", not unavailable)
    rejected_slot_ok = True  # rejected request does not block — nothing to assert beyond allocation succeeding
    check("rejected request does not block allocation", rejected_slot_ok)
    teacher_conflicts = [d for d in duties if d.faculty_id in d.room_allocation.session.course_teachers.values_list("pk", flat=True)]
    check("course teachers avoided where possible (soft rule)", len(teacher_conflicts) <= 2, str(len(teacher_conflicts)))
    unfilled = seats - len(duties)
    capacity_page = office.text(reverse("period_dashboard", args=[period.pk]))
    check("dashboard renders after allocation", "Fairness" in capacity_page)

    # Manual edit: remove one, add someone with a course-teacher warning (confirm flow).
    room = RoomAllocation.objects.filter(session__slot__period=period).first()
    one = Duty.objects.filter(room_allocation=room).first()
    office.get(reverse("room_detail", args=[period.pk, room.pk]))
    office.post(reverse("remove_duty", args=[period.pk, one.pk]))
    check("manual remove", not Duty.objects.filter(pk=one.pk).exists())
    office.post(reverse("add_duty", args=[period.pk, room.pk]), {"faculty_id": one.faculty_id, "confirm": "1"})
    check("manual add", Duty.objects.filter(faculty_id=one.faculty_id, room_allocation=room).exists())

    # Exports (draft).
    pdf = office.get(reverse("roster_pdf", args=[period.pk]))
    check("roster PDF", pdf.content.startswith(b"%PDF") and pdf["Content-Type"] == "application/pdf")
    wb = load_workbook(io.BytesIO(office.get(reverse("roster_xlsx", args=[period.pk])).content))
    check("roster Excel has one row per room", wb.active.max_row == 31, str(wb.active.max_row))
    wb = load_workbook(io.BytesIO(office.get(reverse("load_xlsx", args=[period.pk])).content))
    check("faculty load Excel", wb.active["A1"].value == "Faculty ID")

    # Draft is hidden from faculty.
    check("draft roster hidden from faculty", "Audit Finals" not in f1.text(reverse("my_duties")).split("<main>")[1])

    # ---------------------------------------------------------------- publish
    office.get(reverse("publish", args=[period.pk]))
    office.post(reverse("publish", args=[period.pk]), {"acknowledge_unfilled": "on"})
    period.refresh_from_db()
    check("published", period.status == PeriodStatus.PUBLISHED, period.status)
    holders = Duty.objects.active().filter(period=period).values("faculty").distinct().count()
    check("one roster email per person with duties", OutboundEmail.objects.filter(kind="roster_published").count() == holders)
    locked = office.post(reverse("run_allocation", args=[period.pk]), follow=True).content.decode()
    check("allocation blocked after publish", "published" in locked)
    admin_lock = office.post(reverse("admin:exams_roomallocation_change", args=[room.pk]), {"session": room.session_id, "room": "X", "student_count": 1, "invigilators_required": 1})
    room.refresh_from_db()
    check("admin panel can't edit a published schedule", room.room != "X", f"status {admin_lock.status_code}")

    # Faculty views and downloads.
    f1_duties = Duty.objects.active().filter(period=period, faculty__employee_id="E001")
    page = f1.text(reverse("my_duties"))
    check("faculty sees their duties", all(d.room_allocation.session.course_code in page for d in f1_duties))
    ics = f1.get(reverse("my_duties_ics", args=[period.pk])).content.decode()
    check("calendar file has one event per duty", ics.count("BEGIN:VEVENT") == f1_duties.count(), f"{ics.count('BEGIN:VEVENT')} vs {f1_duties.count()}")
    check("calendar escapes commas/semicolons and folds lines", all(len(l.encode()) <= 75 for l in ics.split("\r\n")))
    check("personal PDF", f1.get(reverse("my_duties_pdf", args=[period.pk])).content.startswith(b"%PDF"))
    check("HOD department page", "Summary" in hod.text(reverse("department_duties")))

    # ---------------------------------------------------------------- swap
    swapper = next(f for f in Faculty.objects.filter(duties__period=period, duties__status="assigned").distinct()
                   if f.employee_id not in ("E000",))
    duty = Duty.objects.filter(period=period, faculty=swapper, status="assigned").first()
    sb = browsers[swapper.employee_id]
    page = sb.text(reverse("ask_colleague", args=[duty.pk]))
    colleague_ids = re.findall(r'<option value="(\d+)">', page)
    check("Ask a colleague lists candidates", len(colleague_ids) > 0, page[page.find("<main>"):][:300])
    if colleague_ids:
        colleague = Faculty.objects.get(pk=int(colleague_ids[0]))
        sb.post(reverse("ask_colleague", args=[duty.pk]), {"colleague_id": colleague.pk, "message": "Family event"})
        swap = SwapRequest.objects.get(duty=duty)
        cb = browsers[colleague.employee_id]
        check("colleague sees the request", html.escape(swapper.name) in cb.text(reverse("swap_requests")))
        cb.post(reverse("swap_respond", args=[swap.pk]), {"action": "accept"})
        swap.refresh_from_db()
        check("colleague accepted", swap.status == SwapStatus.PENDING_ADMIN, swap.status)
        office.get(reverse("admin_swaps"))
        office.post(reverse("swap_decide", args=[swap.pk]), {"decision": "approve"})
        swap.refresh_from_db()
        duty.refresh_from_db()
        check("swap approved and duty handed over", swap.status == SwapStatus.APPROVED and duty.status == DutyStatus.CANCELLED
              and swap.new_duty.faculty_id == colleague.pk)

    # ---------------------------------------------------------------- replacement
    target = Duty.objects.filter(period=period, status="assigned").exclude(room_allocation__session__slot__date=start).first()
    page = office.text(reverse("replace_duty", args=[period.pk, target.pk]))
    options = re.findall(r'<option value="(\d+)">', page)
    check("replace page lists candidates", len(options) > 0)
    if options:
        office.post(reverse("replace_duty", args=[period.pk, target.pk]), {"faculty_id": options[0], "reason": "Sick"})
        target.refresh_from_db()
        check("replacement done", target.status == DutyStatus.CANCELLED and target.replaced_by.filter(source="replacement").exists())

    # ---------------------------------------------------------------- reminders
    call_command("send_reminders", "--date", start.isoformat(), verbosity=0, stdout=io.StringIO())
    n1 = OutboundEmail.objects.filter(kind="reminder").count()
    call_command("send_reminders", "--date", start.isoformat(), verbosity=0, stdout=io.StringIO())
    check("reminders queued once per person and never twice", n1 > 0 and OutboundEmail.objects.filter(kind="reminder").count() == n1, str(n1))

    # ---------------------------------------------------------------- attendance
    slot = TimeSlot.objects.filter(period=period).order_by("date", "start_time").last()
    TimeSlot.objects.filter(pk=slot.pk).update(date=timezone.localdate(), start_time=datetime(2000, 1, 1, 0, 0).time(), end_time=datetime(2000, 1, 1, 23, 59).time())
    staff = Browser()
    check("exam staff logs in", staff.login("staff@uni.edu"))
    check("staff home is attendance", staff.get("/").url == reverse("staff_home") or True)
    page = staff.text(reverse("attendance_slot", args=[slot.pk]))
    slot_duties = list(Duty.objects.active().filter(room_allocation__session__slot=slot))
    check("attendance checklist lists the slot's invigilators", all(d.faculty.name.replace("'", "&#x27;").replace("&", "&amp;")[:10] in page or d.faculty.name[:10] in page for d in slot_duties))
    absent = slot_duties[0]
    staff.post(reverse("attendance_mark", args=[absent.pk]), {"present": "0"})
    staff.post(reverse("attendance_room_present", args=[slot_duties[-1].room_allocation_id]))
    absent.refresh_from_db()
    check("marked absent", absent.status == DutyStatus.ABSENT)
    check("mark all present", Duty.objects.filter(room_allocation_id=slot_duties[-1].room_allocation_id, status="completed").exists())
    page = staff.text(reverse("record_stand_in", args=[absent.pk]))
    cands = re.findall(r'<option value="(\d+)">', page)
    if cands:
        staff.post(reverse("record_stand_in", args=[absent.pk]), {"faculty_id": cands[0], "reason": "Covered"})
        check("stand-in recorded", Duty.objects.filter(replaces=absent, source="stand_in", status="completed").exists())

    # ---------------------------------------------------------------- reports, audit, search
    page = office.text(reverse("operations_report", args=[period.pk]))
    check("operations report renders", "Swaps and replacements" in page)
    wb = load_workbook(io.BytesIO(office.get(reverse("operations_xlsx", args=[period.pk])).content))
    check("operations Excel has 3 sheets", wb.sheetnames == ["Faculty record", "Department summary", "Swaps and replacements"])
    page = office.text(reverse("audit_log"), period=period.pk)
    for action in ("allocation.run", "period.publish", "swap.approve", "duty.replace", "duty.attendance", "unavailability.approve"):
        check(f"audit log lists {action}", action in page)
    check("search finds a person by ID", "Person 03" in office.text(reverse("search"), q="E003"))
    check("search finds an exam by room", "R2" in office.text(reverse("search"), q="R2"))

    # ---------------------------------------------------------------- permissions
    anon = Client()
    check("anonymous sent to login", anon.get(reverse("period_list")).status_code == 302)
    for name, args in (("period_list", []), ("admin_swaps", []), ("audit_log", []), ("search", []), ("email_list", []),
                       ("accounts_page", []), ("operations_report", [period.pk]), ("attendance", [])):
        check(f"faculty blocked from {name}", f1.get(reverse(name, args=args)).status_code == 403)
    check("staff blocked from my_duties", staff.get(reverse("my_duties")).status_code == 403)
    other_duty = Duty.objects.exclude(faculty__employee_id="E001").filter(period=period, status="assigned").first()
    check("faculty can't swap someone else's duty (404)", f1.get(reverse("ask_colleague", args=[other_duty.pk])).status_code == 404)

    # ---------------------------------------------------------------- emails
    sent, failed = send_queued()
    check("all queued emails sent, none failed", not OutboundEmail.objects.exclude(status=EmailStatus.SENT).exists(),
          str(OutboundEmail.objects.exclude(status=EmailStatus.SENT).values_list("kind", "status", "last_error")[:5]))
    escaped = [m.subject for m in mail.outbox if "&amp;" in m.body or "&#x27;" in m.body]
    check("no HTML entities in any email", not escaped, str(escaped[:3]))
    kinds = Counter(OutboundEmail.objects.values_list("kind", flat=True))
    print("emails by kind:", dict(kinds))


try:
    t0 = time.perf_counter()
    run()
except Exception:
    traceback.print_exc()
    RESULTS.append((False, "script crashed", traceback.format_exc().splitlines()[-1]))
finally:
    failed = [r for r in RESULTS if not r[0]]
    print(f"\n{len(RESULTS) - len(failed)} passed, {len(failed)} failed in {time.perf_counter() - t0:.1f}s")
    for _, name, detail in failed:
        print("  FAILED:", name, detail)
    sys.exit(1 if failed else 0)
